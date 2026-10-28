import io
import csv
import datetime

import qrcode

from django.shortcuts import render, get_object_or_404, redirect
from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin, UserPassesTestMixin
from django.views.generic import ListView, DetailView, CreateView, UpdateView, DeleteView
from django.urls import reverse_lazy, reverse
from django.contrib import messages
from django.utils import timezone
from django.http import HttpResponse
from django.db.models import Q
from django.core.files.base import ContentFile
from django.conf import settings

from .models import Activity, Registration, CheckIn
from .forms import ActivityForm


def _get_client_ip(request):
    x_forwarded_for = request.META.get('HTTP_X_FORWARDED_FOR')
    if x_forwarded_for:
        return x_forwarded_for.split(',')[0].strip()
    return request.META.get('REMOTE_ADDR')


def _generate_qr(activity):
    """Generate QR code image for the activity check-in URL and save it."""
    checkin_url = f"/activities/{activity.pk}/checkin/{activity.qr_token}/"
    qr = qrcode.QRCode(
        version=2,
        error_correction=qrcode.constants.ERROR_CORRECT_H,
        box_size=10,
        border=4,
    )
    qr.add_data(checkin_url)
    qr.make(fit=True)
    img = qr.make_image(fill_color="#1a1a2e", back_color="white")
    buffer = io.BytesIO()
    img.save(buffer, format='PNG')
    filename = f"qr_{activity.pk}_{activity.qr_token}.png"
    activity.qr_code.save(filename, ContentFile(buffer.getvalue()), save=True)


# --- Public activity list ---

class ActivityListView(ListView):
    model = Activity
    template_name = 'activities/activity_list.html'
    context_object_name = 'activities'
    paginate_by = 9

    def get_queryset(self):
        qs = Activity.objects.exclude(status='draft').order_by('-created_at')
        q = self.request.GET.get('q', '').strip()
        status = self.request.GET.get('status', '').strip()
        if q:
            qs = qs.filter(Q(title__icontains=q) | Q(location__icontains=q))
        if status:
            qs = qs.filter(status=status)
        return qs

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx['q'] = self.request.GET.get('q', '')
        ctx['status_filter'] = self.request.GET.get('status', '')
        return ctx


# --- Organizer's own activities ---

class OrganizedActivityListView(LoginRequiredMixin, ListView):
    model = Activity
    template_name = 'activities/organized_activity_list.html'
    context_object_name = 'activities'

    def get_queryset(self):
        user = self.request.user
        if user.role == 'admin' or user.is_superuser:
            return Activity.objects.all().order_by('-created_at')
        return Activity.objects.filter(organizer=user).order_by('-created_at')


# --- Activity detail ---

class ActivityDetailView(DetailView):
    model = Activity
    template_name = 'activities/activity_detail.html'

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        user = self.request.user
        if user.is_authenticated:
            reg = Registration.objects.filter(
                activity=self.object, user=user, status='registered'
            ).first()
            ctx['is_registered'] = reg is not None
            ctx['registration'] = reg
            ctx['has_checkin'] = hasattr(reg, 'checkin') if reg else False
        ctx['checkin_count'] = self.object.get_checkin_count()
        return ctx


# --- Create / Update / Delete ---

class ActivityCreateView(LoginRequiredMixin, CreateView):
    model = Activity
    form_class = ActivityForm
    template_name = 'activities/activity_form.html'

    def dispatch(self, request, *args, **kwargs):
        if request.user.is_authenticated and request.user.role not in ('organizer', 'admin') and not request.user.is_superuser:
            messages.error(request, '只有活动组织者或管理员才能发布活动。')
            return redirect('activities:list')
        return super().dispatch(request, *args, **kwargs)

    def get_success_url(self):
        return reverse('activities:organized_list')

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        if self.request.method in ('POST', 'PUT'):
            kwargs['files'] = self.request.FILES
        return kwargs

    def form_valid(self, form):
        form.instance.organizer = self.request.user
        response = super().form_valid(form)
        _generate_qr(self.object)
        messages.success(self.request, f'活动「{self.object.title}」创建成功，二维码已自动生成！')
        return response


class ActivityUpdateView(LoginRequiredMixin, UserPassesTestMixin, UpdateView):
    model = Activity
    form_class = ActivityForm
    template_name = 'activities/activity_form.html'

    def test_func(self):
        activity = self.get_object()
        return self.request.user == activity.organizer or self.request.user.is_superuser

    def get_success_url(self):
        return reverse('activities:organized_list')

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        if self.request.method in ('POST', 'PUT'):
            kwargs['files'] = self.request.FILES
        return kwargs

    def form_valid(self, form):
        messages.success(self.request, '活动信息已更新！')
        return super().form_valid(form)


class ActivityDeleteView(LoginRequiredMixin, UserPassesTestMixin, DeleteView):
    model = Activity
    template_name = 'activities/activity_confirm_delete.html'
    success_url = reverse_lazy('activities:organized_list')

    def test_func(self):
        activity = self.get_object()
        return self.request.user == activity.organizer or self.request.user.is_superuser


# --- Registration ---

@login_required
def register_activity(request, pk):
    activity = get_object_or_404(Activity, pk=pk)
    if activity.status != 'published':
        messages.error(request, '该活动目前不接受报名。')
        return redirect('activities:detail', pk=pk)
    if timezone.now() > activity.reg_deadline:
        messages.error(request, '报名已截止。')
        return redirect('activities:detail', pk=pk)
    if activity.max_participants > 0 and activity.current_participants >= activity.max_participants:
        messages.error(request, '报名人数已满，无法报名。')
        return redirect('activities:detail', pk=pk)
    reg, created = Registration.objects.get_or_create(activity=activity, user=request.user)
    if not created and reg.status == 'registered':
        messages.warning(request, '您已经报名该活动，请勿重复报名。')
    else:
        reg.status = 'registered'
        reg.save()
        activity.current_participants = activity.registrations.filter(status='registered').count()
        activity.save(update_fields=['current_participants'])
        messages.success(request, '报名成功！请在活动现场扫描二维码完成签到。')
    return redirect('activities:detail', pk=pk)


@login_required
def cancel_registration(request, pk):
    activity = get_object_or_404(Activity, pk=pk)
    reg = Registration.objects.filter(activity=activity, user=request.user, status='registered').first()
    if reg:
        if hasattr(reg, 'checkin'):
            messages.error(request, '您已签到，无法取消报名。')
            return redirect('activities:detail', pk=pk)
        reg.status = 'cancelled'
        reg.save()
        activity.current_participants = activity.registrations.filter(status='registered').count()
        activity.save(update_fields=['current_participants'])
        messages.success(request, '已成功取消报名。')
    else:
        messages.error(request, '未找到您的报名记录。')
    return redirect('activities:detail', pk=pk)


# --- QR Code regenerate ---

@login_required
def generate_qr(request, pk):
    activity = get_object_or_404(Activity, pk=pk)
    if request.user != activity.organizer and not request.user.is_superuser:
        messages.error(request, '权限不足。')
        return redirect('activities:detail', pk=pk)
    _generate_qr(activity)
    messages.success(request, '二维码已重新生成！')
    return redirect('activities:detail', pk=pk)


# --- Check-in via QR image upload ---

@login_required
def upload_checkin(request, pk):
    """Student uploads a QR code image to check in."""
    activity = get_object_or_404(Activity, pk=pk)
    if request.method != 'POST':
        return redirect('activities:detail', pk=pk)

    # 验证报名状态
    reg = Registration.objects.filter(activity=activity, user=request.user, status='registered').first()
    if not reg:
        messages.error(request, '您尚未报名该活动，无法签到。')
        return redirect('activities:detail', pk=pk)
    if hasattr(reg, 'checkin'):
        messages.warning(request, '您已经完成签到，请勿重复操作。')
        return redirect('activities:detail', pk=pk)

    qr_image = request.FILES.get('qr_image')
    if not qr_image:
        messages.error(request, '请选择二维码图片文件。')
        return redirect('activities:detail', pk=pk)

    # 解码二维码
    try:
        from PIL import Image
        from pyzbar.pyzbar import decode
        img = Image.open(qr_image)
        results = decode(img)
    except Exception:
        messages.error(request, '图片读取失败，请上传有效的图片文件。')
        return redirect('activities:detail', pk=pk)

    if not results:
        messages.error(request, '未能识别二维码，请确保图片清晰完整。')
        return redirect('activities:detail', pk=pk)

    # 提取二维码内容并验证 token
    qr_data = results[0].data.decode('utf-8')
    expected_token = str(activity.qr_token)
    if expected_token not in qr_data:
        messages.error(request, '二维码不匹配，请使用本活动的签到二维码。')
        return redirect('activities:detail', pk=pk)

    # 验证签到时间窗口
    now = timezone.now()
    window = getattr(settings, 'SIGNIN_WINDOW_MINUTES', 30)
    if activity.status not in ('published', 'ongoing'):
        messages.error(request, '活动当前状态不允许签到。')
        return redirect('activities:detail', pk=pk)
    if now < activity.start_time - datetime.timedelta(minutes=window):
        messages.error(request, f'签到尚未开放，活动开始前 {window} 分钟方可签到。')
        return redirect('activities:detail', pk=pk)
    if now > activity.end_time:
        messages.error(request, '活动已结束，签到通道已关闭。')
        return redirect('activities:detail', pk=pk)

    status = 'normal'
    CheckIn.objects.create(
        registration=reg,
        status=status,
        ip_address=_get_client_ip(request),
        device_info=request.META.get('HTTP_USER_AGENT', '')[:512],
    )
    if status == 'normal':
        messages.success(request, '签到成功！')
    else:
        messages.success(request, '签到成功（迟到）！')
    return redirect('activities:detail', pk=pk)


# --- Check-in via QR scan ---

def checkin_view(request, pk, token):
    activity = get_object_or_404(Activity, pk=pk)
    if str(activity.qr_token) != str(token):
        return render(request, 'activities/checkin_result.html', {
            'success': False, 'message': '二维码无效或已过期，请联系组织者。', 'activity': activity
        })
    if not request.user.is_authenticated:
        return redirect(f"/users/login/?next=/activities/{pk}/checkin/{token}/")
    now = timezone.now()
    window = getattr(settings, 'SIGNIN_WINDOW_MINUTES', 30)
    if activity.status not in ('published', 'ongoing'):
        return render(request, 'activities/checkin_result.html', {
            'success': False, 'message': f'活动当前状态不允许签到。', 'activity': activity
        })
    if now < activity.start_time - datetime.timedelta(minutes=window):
        return render(request, 'activities/checkin_result.html', {
            'success': False, 'message': f'签到尚未开放，活动开始前 {window} 分钟方可签到。', 'activity': activity
        })
    if now > activity.end_time:
        return render(request, 'activities/checkin_result.html', {
            'success': False, 'message': '活动已结束，签到通道已关闭。', 'activity': activity
        })
    reg = Registration.objects.filter(activity=activity, user=request.user, status='registered').first()
    if not reg:
        return render(request, 'activities/checkin_result.html', {
            'success': False, 'message': '您尚未报名该活动，无法签到。', 'activity': activity
        })
    if hasattr(reg, 'checkin'):
        return render(request, 'activities/checkin_result.html', {
            'success': False, 'message': '您已经完成签到，请勿重复操作。',
            'activity': activity, 'checkin': reg.checkin
        })
    status = 'normal'
    checkin = CheckIn.objects.create(
        registration=reg,
        status=status,
        ip_address=_get_client_ip(request),
        device_info=request.META.get('HTTP_USER_AGENT', '')[:512],
    )
    return render(request, 'activities/checkin_result.html', {
        'success': True,
        'message': '签到成功！' if status == 'normal' else '签到成功（迟到）！',
        'activity': activity, 'checkin': checkin,
    })


# --- Statistics ---

@login_required
def activity_stats(request, pk):
    activity = get_object_or_404(Activity, pk=pk)
    if request.user != activity.organizer and request.user.role != 'admin' and not request.user.is_superuser:
        messages.error(request, '权限不足。')
        return redirect('activities:detail', pk=pk)
    registrations = Registration.objects.filter(activity=activity).select_related('user')
    checkins = CheckIn.objects.filter(registration__activity=activity).select_related('registration__user')
    reg_count = registrations.filter(status='registered').count()
    cancel_count = registrations.filter(status='cancelled').count()
    checkin_count = checkins.count()
    normal_count = checkins.filter(status='normal').count()
    late_count = checkins.filter(status='late').count()
    no_show = max(reg_count - checkin_count, 0)
    context = {
        'activity': activity,
        'registrations': registrations.filter(status='registered'),
        'checkins': checkins,
        'reg_count': reg_count,
        'cancel_count': cancel_count,
        'checkin_count': checkin_count,
        'normal_count': normal_count,
        'late_count': late_count,
        'no_show': no_show,
    }
    return render(request, 'activities/activity_stats.html', context)


# --- Export CSV ---

@login_required
def export_registrations_csv(request, pk):
    activity = get_object_or_404(Activity, pk=pk)
    if request.user != activity.organizer and request.user.role != 'admin' and not request.user.is_superuser:
        messages.error(request, '权限不足。')
        return redirect('activities:detail', pk=pk)
    response = HttpResponse(content_type='text/csv; charset=utf-8-sig')
    response['Content-Disposition'] = f'attachment; filename="registrations_{pk}.csv"'
    writer = csv.writer(response)
    writer.writerow(['姓名', '用户名', '学号/工号', '手机号', '报名时间', '报名状态', '是否签到', '签到时间', '签到状态'])
    for reg in Registration.objects.filter(activity=activity).select_related('user'):
        checkin = getattr(reg, 'checkin', None)
        writer.writerow([
            reg.user.real_name or '',
            reg.user.username,
            reg.user.student_id or '',
            reg.user.phone or '',
            reg.registration_time.strftime('%Y-%m-%d %H:%M') if reg.registration_time else '',
            reg.get_status_display(),
            '是' if checkin else '否',
            checkin.check_in_time.strftime('%Y-%m-%d %H:%M') if checkin else '',
            checkin.get_status_display() if checkin else '',
        ])
    return response


# --- Export Excel ---

@login_required
def export_registrations_excel(request, pk):
    activity = get_object_or_404(Activity, pk=pk)
    if request.user != activity.organizer and request.user.role != 'admin' and not request.user.is_superuser:
        messages.error(request, '权限不足。')
        return redirect('activities:detail', pk=pk)
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '报名签到数据'
    headers = ['姓名', '用户名', '学号/工号', '手机号', '报名时间', '报名状态', '是否签到', '签到时间', '签到状态']
    header_fill = PatternFill(start_color='1a1a2e', end_color='1a1a2e', fill_type='solid')
    header_font = Font(bold=True, color='FFFFFF')
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')
        ws.column_dimensions[chr(64 + col)].width = 18
    for row_idx, reg in enumerate(Registration.objects.filter(activity=activity).select_related('user'), 2):
        checkin = getattr(reg, 'checkin', None)
        row_data = [
            reg.user.real_name or '',
            reg.user.username,
            reg.user.student_id or '',
            reg.user.phone or '',
            reg.registration_time.strftime('%Y-%m-%d %H:%M') if reg.registration_time else '',
            reg.get_status_display(),
            '是' if checkin else '否',
            checkin.check_in_time.strftime('%Y-%m-%d %H:%M') if checkin else '',
            checkin.get_status_display() if checkin else '',
        ]
        for col, val in enumerate(row_data, 1):
            ws.cell(row=row_idx, column=col, value=val)
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="registrations_{pk}.xlsx"'
    return response
